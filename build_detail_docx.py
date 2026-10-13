# -*- coding: utf-8 -*-
"""生成《博客详情页测试报告》Word 文档

用例数据从 build_detail_xlsx.py 导入，保证 Word 与 Excel 统计口径一致。
"""
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

from build_detail_xlsx import CASES

OUT = r"D:\study\AI\deepseek-day\博客详情页_测试报告.docx"
EA = "微软雅黑"
doc = Document()

style = doc.styles["Normal"]
style.font.name = "Calibri"
style.font.size = Pt(10.5)
style.element.rPr.rFonts.set(qn("w:eastAsia"), EA)
for sec in doc.sections:
    sec.top_margin = Cm(2.2)
    sec.bottom_margin = Cm(2.2)
    sec.left_margin = Cm(2.4)
    sec.right_margin = Cm(2.4)


def set_ea(run, name=EA):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)


def para(text, size=10.5, bold=False, align=None, space_after=6, color=None, style_name=None):
    p = doc.add_paragraph(style=style_name)
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.bold = bold
    set_ea(r)
    if color:
        r.font.color.rgb = RGBColor(*color)
    if align:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    return p


def heading(text, level):
    h = doc.add_heading(level=level)
    r = h.add_run(text)
    set_ea(r)
    if level == 1:
        r.font.size = Pt(15)
        r.font.color.rgb = RGBColor(0x1F, 0x3B, 0x73)
    elif level == 2:
        r.font.size = Pt(12.5)
        r.font.color.rgb = RGBColor(0x2F, 0x55, 0x97)
    return h


def make_table(headers, rows, widths=None, header_bg="2F5597"):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = t.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = ""
        p = hdr[i].paragraphs[0]
        r = p.add_run(h)
        r.bold = True
        r.font.size = Pt(10)
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        set_ea(r)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        shd = OxmlElement("w:shd")
        shd.set(qn("w:fill"), header_bg)
        hdr[i]._tc.get_or_add_tcPr().append(shd)
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""
            p = cells[i].paragraphs[0]
            r = p.add_run(str(v))
            r.font.size = Pt(9.5)
            set_ea(r)
            p.paragraph_format.space_after = Pt(2)
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Cm(w)
    return t


# ==================== 封面 ====================
para("博客系统 详情页测试报告", size=22, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=6)
para("Blog Detail Page Test Report", size=11, align=WD_ALIGN_PARAGRAPH.CENTER,
     space_after=18, color=(0x60, 0x60, 0x60))

make_table(
    ["项目", "内容"],
    [
        ("被测系统", "博客系统（我的博客系统）"),
        ("被测页面", "http://127.0.0.1:9090/blog_detail.html?blogId=32213"),
        ("页面标题", "博客详情页"),
        ("核心功能", "展示博客标题/日期/正文与作者信息；按权限渲染\"编辑/删除\"按钮"),
        ("关键接口", "GET /blog/getBlogDetail、GET /user/getAuthorInfo、"
                     "POST /blog/delete、POST /blog/update"),
        ("测试类型", "功能测试、安全测试、健壮性测试、兼容性测试"),
        ("测试方式", "Playwright 自动化（Chromium，串行执行）"),
        ("测试工具", "Playwright 1.63.0 + TypeScript"),
        ("测试账号", "zhangsan / 123456"),
        ("测试时间", "2026-10-13"),
        ("报告版本", "V1.0"),
    ],
    widths=[4.0, 12.0],
)
doc.add_paragraph()

# ==================== 一、测试结论 ====================
heading("一、测试结论", 1)
para("本次针对博客详情页共设计 38 条测试用例（114 个测试步骤），覆盖访问控制、数据渲染、"
     "无效 blogId、权限与操作按钮、编辑跳转、删除流程、数据一致性、安全与交互兼容九个方面。")
para("自动化断言层面 38 次执行全部通过，并经过 4 轮重复执行验证稳定性"
     "（152 次执行零失败、零不稳定）。访问控制、删除流程、XSS 防护均表现良好。")
para("测试过程中确认了 7 个缺陷，其中 1 个为高危：作者信息接口泄露密码哈希。"
     "另有 1 项跨用户越权风险因环境限制**未能实证**，已在报告中明确标注。", space_after=10)

para("核心结论：", bold=True, space_after=4)
for line in [
    "① 访问控制良好：未登录访问详情页会被 401 守卫跳转登录页，"
    "详情/作者/删除三个接口对无 token 请求均返回 401，防护有效。",
    "② 删除流程正确：删除有确认框（\"确定删除?\")，取消不会发起请求，"
    "确认后删除成功并跳转列表页；实测删除自建测试博客后原有数据完好。",
    "③ XSS 防护良好：标题使用 .text() 赋值、正文经 markdownToHTML 渲染，"
    "注入的 <img onerror> 与 <script> 均未执行。",
    "④ 安全缺陷：/user/getAuthorInfo 直接返回 password 密码哈希（高危），"
    "与列表页 getUserInfo 属同类问题，说明接口层普遍未过滤敏感字段。",
    "⑤ 权限字段不可信：loginUser 在列表接口为 false、详情接口为 true，"
    "而该字段正是\"编辑/删除\"按钮的显示依据。",
    "⑥ 接口健壮性：更新/删除不存在的 id 均返回 SUCCESS/data:true，误报成功。",
]:
    p = doc.add_paragraph(line)
    p.paragraph_format.left_indent = Cm(0.6)
    p.paragraph_format.space_after = Pt(3)
    for r in p.runs:
        r.font.size = Pt(10.5)
        set_ea(r)

# ==================== 二、测试环境与数据安全 ====================
heading("二、测试环境与数据安全策略", 1)
make_table(
    ["类别", "项目", "说明"],
    [
        ("被测环境", "服务地址", "http://127.0.0.1:9090"),
        ("被测环境", "详情页", "/blog_detail.html?blogId=32213"),
        ("被测环境", "关联页面", "blog_update.html（更新页，由\"编辑\"按钮跳转）"),
        ("运行环境", "操作系统", "Windows"),
        ("运行环境", "浏览器内核", "Chromium 1243（headless）"),
        ("运行环境", "自动化框架", "Playwright 1.63.0 + TypeScript"),
        ("测试数据", "有效账号", "zhangsan / 123456"),
        ("测试数据", "只读样本", "环境中原有博客 blogId=32213"),
    ],
    widths=[3.0, 3.5, 9.5],
)
doc.add_paragraph()

heading("2.1 数据安全策略（重要）", 2)
para("详情页带\"删除\"操作，且更新/删除接口不可逆。为在验证破坏性功能的同时保证"
     "用户原有数据零风险，本次采用如下策略：", space_after=6)
for line in [
    "只读用例使用环境中原有博客 32213，仅读取、绝不修改。",
    "涉及删除与更新的用例，先通过 /blog/add 创建**专属测试博客**"
    "（标题带 __E2E_DETAIL__ 标记），只对自建数据执行破坏性操作。",
    "每个涉及自建数据的用例都在 finally 中清理测试博客；测试结束后校验环境"
    "仅剩原有博客 32213，标题与内容与测试前完全一致。",
    "越权与异常用例使用**不存在的 id**（如 99999999），不触碰任何真实数据。",
    "专门设置用例 F4 验证\"删除测试博客后原有博客完好、环境恢复\"，实测通过。",
]:
    p = doc.add_paragraph(line, style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    for r in p.runs:
        r.font.size = Pt(10.5)
        set_ea(r)

heading("2.2 测试数据生命周期", 2)
para("实际执行中创建并清理了 6 篇测试博客（用于权限按钮、编辑跳转、删除确认、取消删除、"
     "真实删除、环境恢复等用例）。每篇均在使用后立即删除，最终环境中仅保留原有博客 "
     "blogId=32213，无任何测试数据残留。", space_after=8)

# ==================== 三、执行统计 ====================
heading("三、测试执行统计", 1)

heading("3.1 总体统计", 2)
defect_cases = [c[0] for c in CASES if c[8].startswith("发现缺陷")]
make_table(
    ["统计项", "数值", "说明"],
    [
        ("用例条数", str(len(CASES)), "Excel 用例文档中的用例条数"),
        ("测试步骤数", str(sum(len(c[7]) for c in CASES)), "每条用例按步骤逐行展开"),
        ("自动化执行次数", str(len(CASES)), "1 条用例 = 1 次执行"),
        ("执行通过", str(len(CASES)), "断言层面全部通过"),
        ("执行失败", "0", "无断言失败"),
        ("通过率", "100%", f"{len(CASES)} / {len(CASES)}"),
        ("稳定性验证", "152 次通过", f"{len(CASES)} 条 × 重复 4 轮，零失败、零不稳定"),
        ("发现缺陷", "7", "1 高 / 3 中 / 3 低"),
        ("结论为\"发现缺陷\"的用例", str(len(defect_cases)), "、".join(defect_cases)),
        ("数据残留", "无", "环境仅剩原有博客 32213"),
        ("未能结论项", "1 项", "跨用户越权（环境无注册接口，无法创建第二用户）"),
    ],
    widths=[5.0, 2.5, 8.5],
)
doc.add_paragraph()

heading("3.2 模块分布", 2)
MODULE_ORDER = ["A 访问控制", "B 数据渲染", "C 无效与异常 blogId", "D 权限与操作按钮",
                "E 编辑跳转", "F 删除流程", "G 数据一致性", "H 安全", "I 交互与兼容"]
_mod_rows = []
for _m in MODULE_ORDER:
    _sub = [c for c in CASES if c[1] == _m]
    _mod_rows.append((_m, str(len(_sub)), str(sum(len(c[7]) for c in _sub)),
                      str(len(_sub)), "0", "100%"))
_mod_rows.append(("合计", str(len(CASES)), str(sum(len(c[7]) for c in CASES)),
                  str(len(CASES)), "0", "100%"))
make_table(
    ["模块", "用例数", "步骤数", "通过", "失败", "通过率"],
    _mod_rows,
    widths=[4.4, 2.0, 2.0, 1.8, 1.8, 2.0],
)
doc.add_paragraph()

heading("3.3 缺陷严重级别分布", 2)
make_table(
    ["严重级别", "数量", "占比", "缺陷编号"],
    [
        ("高", "1", "14.3%", "BUG-D01"),
        ("中", "3", "42.9%", "BUG-D02、BUG-D03、BUG-D04"),
        ("低", "3", "42.9%", "BUG-D05、BUG-D06、BUG-D07"),
        ("合计", "7", "100%", "—"),
    ],
    widths=[3.0, 2.0, 2.5, 8.5],
)
doc.add_paragraph()

# ==================== 四、缺陷明细 ====================
heading("四、缺陷明细", 1)

DEFECTS = [
    dict(no="BUG-D01", sev="高", cat="安全", title="/user/getAuthorInfo 响应体泄露密码哈希",
         desc="获取博客作者信息的接口在 data 中直接返回 password 字段，值为用户密码哈希。"
              "详情页仅需 userName 与 githubUrl，完全不需要密码字段。",
         steps="1）登录后以有效 token 调用 /user/getAuthorInfo?blogId=32213；\n"
               "2）查看响应 data 的全部字段。",
         actual="响应字段为 [id, userName, password, githubUrl, deleteFlag, createTime, updateTime]，"
                "其中 password 为密码哈希值。",
         impact="任何登录用户都可通过该接口获取任意博客作者的密码哈希；结合列表页的高危 XSS"
                "（BUG-L01）可批量窃取，风险叠加。该缺陷与列表页 BUG-L02（getUserInfo）为同类问题，"
                "说明接口层普遍未做敏感字段过滤。",
         fix="统一在接口返回前剔除 password 等敏感字段，改用 VO/DTO 只输出前端需要的字段。",
         ref="H3"),
    dict(no="BUG-D02", sev="中", cat="数据一致性", title="loginUser 字段在列表与详情接口间矛盾",
         desc="同一篇博客，列表接口 /blog/getList 返回 loginUser=false，"
              "详情接口 /blog/getBlogDetail 返回 loginUser=true。该字段用于决定是否显示"
              "\"编辑/删除\"按钮，取值矛盾意味着字段语义不可信。",
         steps="1）以 zhangsan 登录；\n2）调用 /blog/getList 查看博客 32213 的 loginUser；\n"
               "3）调用 /blog/getBlogDetail?blogId=32213 查看同一字段。",
         actual="列表接口 loginUser=false，详情接口 loginUser=true，两者矛盾。",
         impact="权限展示所依赖的字段不可信：若某接口把他人博客算成 true，会在前端暴露"
                "编辑/删除入口（后端是否兜底校验未确认，见 BUG-D03）。",
         fix="统一 loginUser 的计算逻辑并补充自动化断言；权限判断必须在后端强制校验，"
             "不能只依赖前端字段。",
         ref="D3"),
    dict(no="BUG-D03", sev="中", cat="数据完整性", title="更新/删除不存在的 id 被误报为成功",
         desc="/blog/update 与 /blog/delete 对不存在的 blogId 仍返回 code=SUCCESS、data=true，"
              "成功判定不基于数据库影响行数。",
         steps="1）以有效 token 提交 POST /blog/update，id=99999999；\n"
               "2）以有效 token 提交 POST /blog/delete?blogId=99999999；\n3）查看响应。",
         actual="两个接口均返回 {\"code\":\"SUCCESS\",\"data\":true}；"
                "实际未修改/删除任何记录（已验证其他数据未被误改，WHERE 条件正常）。",
         impact="用户会收到\"操作成功\"的假象，但数据并未变化；同时说明接口未校验记录是否存在，"
                "归属校验也极可能缺失，是越权风险的重要信号。",
         fix="更新/删除应基于影响行数判定成功（affectedRows > 0），记录不存在时返回明确失败；"
             "并强制校验 blog.userId 是否等于当前登录用户。",
         ref="D4、H4、H6"),
    dict(no="BUG-D04", sev="中", cat="可用性", title="无效 blogId 错误信息不可诊断且无友好空态",
         desc="参数非法（abc）、记录不存在（99999999）、参数缺失三种情况返回完全相同的"
              "\"内部错误, 请联系管理员\"；页面仅弹该提示，标题与正文均空白，无返回入口。",
         steps="1）分别以 blogId=abc、blogId=99999999、不带 blogId 访问详情接口；\n"
               "2）在页面上以不存在的 id 打开详情页。",
         actual="三种情况接口响应完全相同；页面上标题、日期、正文均为空且无操作按钮，"
                "仅弹出\"内部错误, 请联系管理员\"。",
         impact="用户与开发者都无法区分是链接错误、文章已删除还是系统故障；"
                "用户停留在空白页面无法继续操作，属可用性与可维护性问题。",
         fix="区分参数非法与记录不存在并返回明确 errMsg；页面在无数据时展示空态提示"
             "（如\"该博客不存在或已被删除\"）并提供返回列表入口。",
         ref="C1、C2、C3"),
    dict(no="BUG-D05", sev="低", cat="数据一致性", title="左侧文章计数为硬编码，与实际不符",
         desc="左侧卡片中\"文章/分类\"的计数为 HTML 静态写死（2 和 1），未随实际数据变化。",
         steps="1）调用 /blog/getList 统计当前账号实际博客数；\n"
               "2）打开详情页读取卡片显示的文章数；\n3）比对两者。",
         actual="卡片显示文章数=2，实际博客数=1，偏差 1。",
         impact="用户看到的统计与实际内容不符。该问题在列表页同样存在（BUG-L03），"
                "属系统性硬编码问题。",
         fix="计数改为动态获取（由接口返回统计数据或按列表长度计算）。",
         ref="G1"),
    dict(no="BUG-D06", sev="低", cat="代码质量", title="详情页存在重复与陈旧代码",
         desc="页面源码中保留了一个无参的 deleteBlog() 函数（且带 confirm 逻辑）、"
              "重复定义的 userUrl 变量，以及被注释掉的 getUserInfo 调用。",
         steps="1）查看 /blog_detail.html 源码；\n2）检查脚本段的函数与变量定义。",
         actual="存在被注释的 deleteBlog() 定义与重复的 userUrl 变量声明；"
                "页面注释中关于 Markdown 渲染方式的说明与实际实现不一致。",
         impact="死代码增加维护成本与误改风险，重复声明的变量可能导致后续修改时行为不符预期。",
         fix="清理注释掉的死代码与重复声明，统一由后端返回值驱动前端逻辑。",
         ref="—"),
    dict(no="BUG-D07", sev="低", cat="兼容性", title="窄屏下出现横向滚动条",
         desc=".container 为固定宽度 800px，未做响应式收缩。在 768px、375px 等窄视口下，"
              "页面宽度超出视口，出现横向滚动条。",
         steps="1）打开详情页；\n2）将视口依次设为 1920、1366、768、375 宽度；\n"
               "3）测量 container 宽度与 document.scrollWidth。",
         actual="container 宽度恒为 800px；视口 768px 时文档宽 800px（溢出 32px），"
                "视口 375px 时文档宽 800px（溢出 425px），均出现横向滚动。",
         impact="移动端或窄窗口下用户需要左右拖动才能阅读正文，体验较差。"
                "该问题同样存在于列表页与编辑页（共用 common.css）。",
         fix="为 .container 设置 max-width: 100% 并配合媒体查询做响应式布局。",
         ref="I3"),
]

SEV_COLOR = {"高": (0xC0, 0x00, 0x00), "中": (0xBF, 0x60, 0x00), "低": (0x40, 0x40, 0x40)}
for i, d in enumerate(DEFECTS, start=1):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(f"{i}. [{d['no']}] {d['title']}")
    r.bold = True
    r.font.size = Pt(11.5)
    r.font.color.rgb = RGBColor(*SEV_COLOR[d["sev"]])
    set_ea(r)
    make_table(
        ["字段", "内容"],
        [
            ("严重级别", d["sev"]),
            ("缺陷类别", d["cat"]),
            ("缺陷描述", d["desc"]),
            ("复现步骤", d["steps"]),
            ("实际结果", d["actual"]),
            ("影响", d["impact"]),
            ("修复建议", d["fix"]),
            ("关联用例", d["ref"]),
            ("状态", "待修复"),
        ],
        widths=[2.8, 13.2],
    )
    doc.add_paragraph()

doc.add_page_break()

# ==================== 五、测试范围 ====================
heading("五、测试范围与用例执行情况", 1)
para("下表汇总各模块用例的执行情况；完整用例（含前置条件、测试数据、逐步操作、"
     "预期结果与实际结果）详见配套 Excel 文档《博客详情页_测试用例.xlsx》。", space_after=8)
make_table(
    ["模块", "用例编号", "用例数", "覆盖内容", "执行情况"],
    [
        ("A 访问控制", "A1-A4", "4",
         "未登录访问页面、详情/作者/删除接口的 401 防护", "全部通过（防护有效）"),
        ("B 数据渲染", "B1-B5", "5",
         "字段一致性、Markdown 渲染、作者卡片、长正文、特殊字符", "全部通过"),
        ("C 无效与异常 blogId", "C1-C5", "5",
         "不存在/非数字/缺参数/负数 id、非 JSON 响应", "确认 1 个缺陷"),
        ("D 权限与操作按钮", "D1-D4", "4",
         "按钮按权限显示、loginUser 字段一致性、后端归属校验风险", "确认 2 个缺陷"),
        ("E 编辑跳转", "E1-E3", "3",
         "编辑按钮跳转、更新页预填、按钮文案", "全部通过"),
        ("F 删除流程", "F1-F4", "4",
         "确认框、取消不删、确认删除、删除后环境恢复", "全部通过（流程正确）"),
        ("G 数据一致性", "G1-G2", "2", "卡片计数、详情与列表字段一致性", "确认 1 个缺陷"),
        ("H 安全", "H1-H6", "6",
         "标题/正文 XSS、作者信息泄露、更新删除不存在的 id", "确认 2 个缺陷"),
        ("I 交互与兼容", "I1-I5", "5",
         "导航、注销、窄屏溢出、静态资源、无效 token", "确认 1 个缺陷"),
    ],
    widths=[3.0, 2.2, 1.6, 6.2, 3.0],
)
doc.add_paragraph()

heading("5.1 已验证正常的关键项", 2)
for line in [
    "访问控制有效：未登录访问详情页会被守卫跳转登录页；详情、作者信息、删除三个接口"
    "对无 token 请求均返回 401；无效 token 同样被拦截。",
    "权限按钮控制正确：接口返回 loginUser=false 时，页面不渲染任何\"编辑/删除\"按钮"
    "（前端按字段正确隐藏）。",
    "删除流程完整正确：点击删除弹出\"确定删除?\"确认框；取消时不发起任何删除请求且博客仍在；"
    "确认后删除成功并跳转列表页，列表中该博客消失。",
    "XSS 防护良好：标题中的 <img onerror> 未产生元素、未执行脚本（使用 .text() 赋值）；"
    "正文中的 <img onerror> 与 <script> 经 markdownToHTML 渲染后也均未执行。",
    "数据渲染正确：标题、日期与接口数据一致；正文渲染为真实 HTML 元素；"
    "作者卡片用户名与 GitHub 链接来自接口；长正文（120 段）不被截断。",
    "编辑跳转正确：跳转 blog_update.html 并带上正确 blogId；更新页正确预填原标题、原正文，"
    "按钮文案为\"更新文章\"。",
    "数据未误伤：缺少 id 的更新不修改任何数据（返回 data=false）；"
    "删除不存在的 id 未误删其他数据；删除自建测试博客后原有博客完好、环境完全恢复。",
    "静态资源正常：本次运行未出现 4xx/5xx 资源。",
]:
    p = doc.add_paragraph(line, style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    for r in p.runs:
        r.font.size = Pt(10.5)
        set_ea(r)

heading("5.2 未能结论的风险项（需后端配合确认）", 2)
para("跨用户越权修改/删除：详情页的\"编辑/删除\"按钮由接口返回的 loginUser 字段驱动，"
     "属纯前端控制。本次**无法创建第二个用户**（系统无注册接口），因此未能实证"
     "\"用户 A 能否修改或删除用户 B 的博客\"。", space_after=6)
para("已确认的相关事实：更新/删除接口对**不存在的 id** 也返回 SUCCESS/data:true，"
     "说明其成功判定不基于数据库影响行数（BUG-D03），归属校验很可能同样缺失。"
     "因此该项属于**高风险待确认项**，建议后端补充校验并单独验证：", space_after=6)
for line in [
    "在 /blog/update 与 /blog/delete 中校验 blog.userId == 当前登录用户 id；",
    "补充跨用户越权的自动化用例（需先提供可注册或可造数据的测试环境）；",
    "为验证提供第二个测试账号，以便覆盖多用户权限场景。",
]:
    p = doc.add_paragraph(line, style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    for r in p.runs:
        r.font.size = Pt(10.5)
        set_ea(r)

heading("5.3 测试限制", 2)
para("① 未覆盖 blog_update.html 的更新功能本身（仅验证了跳转与预填），"
     "更新提交属独立页面，建议单独测试。", space_after=3)
para("② 未验证删除后数据库是物理删除还是逻辑删除（仅有 deleteFlag 字段线索）。", space_after=3)
para("③ 跨用户越权未能实证，原因见 5.2。", space_after=8)

# ==================== 六、风险评估与建议 ====================
heading("六、风险评估与建议", 1)
para("本页最需要关注的是**权限与敏感信息**两方面。BUG-D01（密码哈希泄露）与列表页 XSS "
     "组合后可批量窃取账号；BUG-D02/D03 反映权限判定与写入校验存在系统性薄弱点，"
     "建议按以下优先级处理。", space_after=8)
make_table(
    ["优先级", "建议项", "对应缺陷"],
    [
        ("立即修复", "getAuthorInfo 剔除 password 等敏感字段（与 getUserInfo 一并整改）",
         "BUG-D01"),
        ("立即修复", "更新/删除强制校验 blog.userId 归属，并基于影响行数判定成功",
         "BUG-D03、BUG-D02"),
        ("建议修复", "统一 loginUser 计算逻辑，权限判断在后端强制执行", "BUG-D02"),
        ("建议修复", "区分参数非法与记录不存在，页面补充空态提示与返回入口", "BUG-D04"),
        ("可选优化", "计数改为动态获取", "BUG-D05"),
        ("可选优化", "清理死代码与重复声明；文件级 XSS 复核", "BUG-D06"),
        ("可选优化", "container 增加响应式宽度（三页共用样式一并修复）", "BUG-D07"),
        ("测试补充", "提供第二个测试账号，补齐跨用户越权用例", "5.2"),
    ],
    widths=[2.6, 10.4, 3.0],
)
doc.add_paragraph()

# ==================== 七、附录 ====================
heading("七、附录：测试资产", 1)
make_table(
    ["资产", "路径", "说明"],
    [
        ("测试用例文档", "博客详情页_测试用例.xlsx",
         "38 条用例 / 114 步，含用例说明、用例明细、模块划分、缺陷清单、执行统计 5 个工作表"),
        ("测试报告", "博客详情页_测试报告.docx", "本报告"),
        ("测试规划", "blog_detail_test_plan.md", "测试点、模块与数据安全策略"),
        ("用例脚本", "tests/blog_detail.spec.ts", "Playwright 用例实现，38 条用例"),
        ("页面对象", "pages/DetailPage.ts", "详情页 Page Object（含测试数据创建/清理助手）"),
        ("一键执行", "run-tests-detail.cmd", "免配置执行脚本，自动探测浏览器目录"),
        ("HTML 报告", "playwright-report/index.html", "含失败截图与 trace 回放"),
    ],
    widths=[3.0, 5.4, 7.6],
)
doc.add_paragraph()

para("说明：本报告数据来源为 2026-10-13 的自动化执行结果（Playwright，38 条用例全部通过，"
     "并完成 4 轮 152 次重复稳定性验证）。所有缺陷结论均经实际运行验证，非基于代码静态推断。"
     "测试全程未修改或删除任何原有数据。",
     size=9.5, color=(0x60, 0x60, 0x60))

doc.save(OUT)
print("已生成:", OUT)
print("段落数:", len(doc.paragraphs), "表格数:", len(doc.tables))
