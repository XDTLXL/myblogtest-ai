# -*- coding: utf-8 -*-
"""生成《博客详情页测试用例》—— 规范测试用例格式（步骤级展开）"""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

OUT = r"D:\study\AI\deepseek-day\博客详情页_测试用例.xlsx"

HDR_FILL = PatternFill("solid", fgColor="2F5597")
HDR_FONT = Font(bold=True, color="FFFFFF", size=10.5, name="微软雅黑")
TITLE_FONT = Font(bold=True, size=14, name="微软雅黑")
FONT = Font(size=10, name="微软雅黑")
BOLD = Font(bold=True, size=10, name="微软雅黑")
BOLD_RED = Font(bold=True, size=10, name="微软雅黑", color="C00000")
THIN = Side(style="thin", color="B4C6E7")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
WRAP_C = Alignment(wrap_text=True, vertical="center", horizontal="center")
CENTER = Alignment(horizontal="center", vertical="center")
CASE_A = PatternFill("solid", fgColor="FFFFFF")
CASE_B = PatternFill("solid", fgColor="F2F7FC")
SEC_FILL = PatternFill("solid", fgColor="DDEBF7")

LOGIN_PRE = ("1.被测环境 http://127.0.0.1:9090 已启动\n"
             "2.已用 zhangsan/123456 登录并持有有效 user_token")
SAFE = "\n【数据安全】只针对自建测试博客（标题含 __E2E_DETAIL__）操作，用例结束即清理"
READONLY = "\n【只读】使用环境中原有博客 32213，不做任何修改"

CASES = [
    ("A1", "A 访问控制", "未登录访问详情页被守卫拦截", "P0", "场景法",
     "1.未登录（localStorage 无 token）\n2.目标博客 id=32213",
     "访问 /blog_detail.html?blogId=32213",
     [("清空 localStorage 登录态", "token 为空", "token 已清空"),
      ("直接访问详情页 URL", "接口返回 401，error 回调跳转登录页", "跳转 blog_login.html"),
      ("确认最终地址", "URL 为 blog_login.html", "URL 正确，守卫有效")],
     "通过"),
    ("A2", "A 访问控制", "未登录调用详情接口返回 401", "P0", "错误推测",
     "1.接口可访问\n2.不携带 token",
     "GET /blog/getBlogDetail?blogId=32213（无 token）",
     [("不携带 token 调用详情接口", "返回 HTTP 401", "返回 401，拒绝访问")],
     "通过"),
    ("A3", "A 访问控制", "未登录调用作者信息接口返回 401", "P0", "错误推测",
     "1.接口可访问\n2.不携带 token",
     "GET /user/getAuthorInfo?blogId=32213（无 token）",
     [("不携带 token 调用作者信息接口", "返回 HTTP 401", "返回 401，拒绝访问")],
     "通过"),
    ("A4", "A 访问控制", "未登录调用删除接口返回 401 且数据未被删除", "P0", "场景法",
     "1.已登录以获取博客基线\n2.删除请求不带 token",
     "POST /blog/delete?blogId=32213（无 token）" + READONLY,
     [("记录当前博客条数与 id 列表", "作为基线", "基线 1 条，id=32213"),
      ("不携带 token 调用删除接口", "返回 401，写入被拒绝", "返回 401"),
      ("再次统计博客条数", "与基线一致，博客未被删除", "仍为 1 条且 32213 存在，数据完好")],
     "通过"),

    ("B1", "B 数据渲染", "标题、日期、正文与接口数据一致", "P0", "功能测试",
     LOGIN_PRE + READONLY,
     "blogId=32213",
     [("调用详情接口获取期望值", "返回 title、createTime、content", "接口返回成功"),
      ("打开详情页并等待渲染", "页面填充数据", "渲染完成"),
      ("比对标题", "与接口 title 完全一致", "一致"),
      ("比对日期", "包含接口 createTime 的日期部分", "一致")],
     "通过"),
    ("B2", "B 数据渲染", "正文经 Markdown 渲染生成容器", "P0", "功能测试",
     LOGIN_PRE + READONLY,
     "blogId=32213，正文内容\"博客内容\"",
     [("打开详情页并等待渲染", "正文区填充内容", "渲染完成"),
      ("检查是否生成 .markdown-body 容器", "editor.md 应生成该容器", "存在 1 个 markdown-body 容器"),
      ("检查正文是否为真实 HTML 元素", "正文应渲染为 <p> 等元素而非纯文本", "渲染出 <p>博客内容</p>")],
     "通过"),
    ("B3", "B 数据渲染", "作者信息填充左侧卡片", "P1", "功能测试",
     LOGIN_PRE + READONLY,
     "blogId=32213；作者 zhangsan",
     [("调用作者信息接口获取期望值", "返回 userName、githubUrl", "接口返回成功"),
      ("打开详情页等待卡片填充", "卡片显示作者信息", "卡片已填充"),
      ("比对卡片用户名", "等于接口 userName", "显示 zhangsan，一致"),
      ("比对 GitHub 链接", "等于接口 githubUrl", "链接一致")],
     "通过"),
    ("B4", "B 数据渲染", "长正文渲染不被截断", "P2", "边界值分析",
     LOGIN_PRE,
     "拦截详情接口返回 120 段正文（约 2000 字）",
     [("拦截详情接口返回长正文", "拦截生效", "拦截生效"),
      ("打开详情页并等待渲染", "正文完整渲染", "渲染完成"),
      ("检查末尾段落是否存在", "应能读到第 120 段", "读到第 120 段，未截断")],
     "通过"),
    ("B5", "B 数据渲染", "特殊字符标题作为文本渲染", "P1", "错误推测",
     LOGIN_PRE,
     "拦截详情接口返回标题：引号\"单引号'与&符号<>尖括号",
     [("拦截详情接口返回特殊字符标题", "拦截生效", "拦截生效"),
      ("打开详情页并读取标题", "标题原样显示，无乱码", "与输入完全一致"),
      ("检查标题内是否产生子元素", "使用 .text() 赋值，不解析 HTML", "子元素数 0")],
     "通过"),

    ("C1", "C 无效与异常 blogId", "不存在的 blogId：空白页面并提示内部错误", "P1", "错误推测",
     LOGIN_PRE,
     "blogId=99999999（不存在）",
     [("以不存在的 blogId 打开详情页", "页面应给出友好提示或空态", "标题与正文均为空，无操作按钮"),
      ("监听弹窗", "记录提示内容", "弹出\"内部错误, 请联系管理员\""),
      ("评估影响", "无友好空态与返回入口，属缺陷", "确认无友好空态，已记为 BUG-D04")],
     "发现缺陷 BUG-D04"),
    ("C2", "C 无效与异常 blogId", "非数字 blogId 的错误信息与\"不存在\"无法区分", "P1", "错误推测",
     "1.接口可访问\n2.有效 token",
     "对比 blogId=abc 与 blogId=99999999",
     [("以 blogId=abc 调用详情接口", "记录错误信息", "code=FAIL，errMsg=\"内部错误, 请联系管理员\""),
      ("以 blogId=99999999 调用详情接口", "记录错误信息", "与上面完全相同"),
      ("比对两者", "参数非法与记录不存在应可区分", "错误信息完全一致，无法定位原因"),
      ("评估影响", "错误信息不可诊断，属缺陷", "确认无法区分，已记为 BUG-D04")],
     "发现缺陷 BUG-D04"),
    ("C3", "C 无效与异常 blogId", "缺少 blogId 参数时表现一致", "P1", "边界值分析",
     "1.接口可访问\n2.有效 token",
     "GET /blog/getBlogDetail（不带 blogId）",
     [("不携带 blogId 调用详情接口", "记录错误信息", "同样返回\"内部错误, 请联系管理员\""),
      ("与不存在 id 的错误信息比对", "应能区分参数缺失与记录不存在", "三者错误信息完全相同（BUG-D04）")],
     "发现缺陷 BUG-D04"),
    ("C4", "C 无效与异常 blogId", "负数 blogId 失败但不返回 5xx", "P2", "边界值分析",
     "1.接口可访问\n2.有效 token",
     "blogId=-1",
     [("以 blogId=-1 调用详情接口", "返回可读失败，不返回 5xx", "状态码小于 500，code=FAIL")],
     "通过"),
    ("C5", "C 无效与异常 blogId", "详情接口返回非 JSON 时页面不崩溃", "P1", "错误推测",
     LOGIN_PRE,
     "拦截详情接口返回 text/html 内容",
     [("拦截详情接口返回非 JSON", "拦截生效", "拦截生效"),
      ("打开详情页", "页面不白屏、不崩溃", "左侧卡片仍可见，页面正常"),
      ("检查控制台异常", "无未捕获异常", "无异常")],
     "通过"),

    ("D1", "D 权限与操作按钮", "loginUser=true 时渲染编辑与删除按钮", "P0", "功能测试",
     LOGIN_PRE + SAFE,
     "自建测试博客（当前用户所属）",
     [("创建专属测试博客", "创建成功并拿到 id", "创建成功"),
      ("打开该博客详情页", "页面渲染完成", "渲染完成"),
      ("读取操作区按钮", "应渲染\"编辑\"与\"删除\"按钮", "两个按钮均存在"),
      ("清理测试博客", "删除测试数据", "已清理")],
     "通过"),
    ("D2", "D 权限与操作按钮", "loginUser=false 时不渲染任何操作按钮", "P0", "错误推测",
     LOGIN_PRE,
     "拦截详情接口返回 loginUser=false 的他人博客",
     [("拦截详情接口返回 loginUser=false", "拦截生效", "拦截生效"),
      ("打开详情页并读取操作区", "不应渲染任何操作按钮", "按钮数为 0，前端控制正确")],
     "通过"),
    ("D3", "D 权限与操作按钮", "loginUser 字段在列表与详情接口间矛盾", "P1", "错误推测",
     LOGIN_PRE + READONLY,
     "同一博客 32213 分别调用列表与详情接口",
     [("调用列表接口查看该博客", "记录 loginUser 值", "列表接口返回 loginUser=false"),
      ("调用详情接口查看同一博客", "记录 loginUser 值", "详情接口返回 loginUser=true"),
      ("比对两接口取值", "同一博客的同一字段应一致", "取值矛盾，字段语义不可信（BUG-D02）")],
     "发现缺陷 BUG-D02"),
    ("D4", "D 权限与操作按钮", "按钮为前端控制，后端归属校验需确认", "P0", "错误推测",
     LOGIN_PRE,
     "用有效 token 更新不存在的 id=99999999",
     [("以有效 token 更新不存在的 id", "应返回失败", "返回 code=SUCCESS、data=true"),
      ("分析成功判定依据", "应基于影响行数判定", "成功判定不基于影响行数，归属校验存在风险"),
      ("说明限制", "环境无注册接口，无法创建第二用户", "跨用户越权未能实证，已在报告中列为待确认风险")],
     "发现缺陷 BUG-D03"),

    ("E1", "E 编辑跳转", "点击编辑跳转到更新页并带上 blogId", "P1", "功能测试",
     LOGIN_PRE + SAFE,
     "自建测试博客",
     [("创建专属测试博客并打开详情页", "页面渲染出编辑按钮", "按钮存在"),
      ("点击\"编辑\"按钮", "跳转 blog_update.html?blogId=<当前id>", "跳转成功且参数正确"),
      ("清理测试博客", "删除测试数据", "已清理")],
     "通过"),
    ("E2", "E 编辑跳转", "更新页正确预填原博客数据", "P1", "功能测试",
     LOGIN_PRE + SAFE,
     "自建测试博客",
     [("进入 blog_update.html?blogId=<id>", "页面加载并请求原数据", "加载完成"),
      ("读取隐藏字段 #blogId", "等于当前博客 id", "取值正确"),
      ("读取 #title", "为原博客标题", "与原标题一致"),
      ("读取编辑器内容", "为原博客正文", "与原文一致")],
     "通过"),
    ("E3", "E 编辑跳转", "更新页按钮文案为\"更新文章\"", "P2", "界面检查",
     LOGIN_PRE + SAFE,
     "自建测试博客",
     [("进入更新页", "页面加载完成", "加载完成"),
      ("读取 #submit 的 value", "应为\"更新文章\"", "文案正确"),
      ("清理测试博客", "删除测试数据", "已清理")],
     "通过"),

    ("F1", "F 删除流程", "点击删除弹出确认框\"确定删除?\"", "P0", "功能测试",
     LOGIN_PRE + SAFE,
     "自建测试博客",
     [("创建测试博客并打开详情页", "渲染出删除按钮", "按钮存在"),
      ("点击\"删除\"按钮并监听弹窗", "弹出确认对话框", "弹出 confirm 框"),
      ("读取确认文案", "应为\"确定删除?\"", "文案为\"确定删除?\""),
      ("清理测试博客", "删除测试数据", "已清理")],
     "通过"),
    ("F2", "F 删除流程", "取消确认后不删除且不发起删除请求", "P0", "场景法",
     LOGIN_PRE + SAFE,
     "自建测试博客",
     [("创建测试博客并打开详情页", "渲染出删除按钮", "按钮存在"),
      ("监听删除接口请求", "作为判定依据", "已监听"),
      ("点击删除后在确认框中选择取消", "不发起删除请求", "删除请求数为 0"),
      ("检查页面地址与博客状态", "停留在详情页，博客仍存在", "地址未变，博客仍在"),
      ("清理测试博客", "删除测试数据", "已清理")],
     "通过"),
    ("F3", "F 删除流程", "确认删除后跳转列表页且博客消失", "P0", "功能测试",
     LOGIN_PRE + SAFE,
     "自建测试博客",
     [("创建测试博客并打开详情页", "渲染出删除按钮", "按钮存在"),
      ("点击删除并在确认框中选择确定", "发起删除请求并处理结果", "请求成功"),
      ("观察页面跳转", "跳转 blog_list.html", "跳转成功"),
      ("再次查询列表", "该测试博客已消失", "列表中已不存在")],
     "通过"),
    ("F4", "F 删除流程", "删除测试博客后原有博客完好、环境恢复", "P0", "场景法",
     LOGIN_PRE + READONLY,
     "自建测试博客 + 原有博客 32213",
     [("记录删除前的博客 id 基线", "作为比对依据", "基线 1 条：32213"),
      ("创建测试博客", "条数变为基线+1", "条数 +1"),
      ("删除测试博客", "删除成功", "删除成功"),
      ("比对最终 id 列表与基线", "应完全一致，原有博客完好", "一致，32213 完好，环境已恢复")],
     "通过"),

    ("G1", "G 数据一致性", "左侧文章计数为硬编码，与实际不符", "P1", "界面检查",
     LOGIN_PRE + READONLY,
     "blogId=32213；当前账号实际博客数",
     [("调用列表接口统计实际博客数", "记录实际数量", "实际 1 条"),
      ("打开详情页读取卡片计数", "记录卡片显示数量", "卡片显示文章数=2"),
      ("比对实际值与显示值", "应一致", "不一致（偏差 1），卡片为 HTML 硬编码（BUG-D05）")],
     "发现缺陷 BUG-D05"),
    ("G2", "G 数据一致性", "详情字段与列表接口一致", "P1", "功能测试",
     LOGIN_PRE + READONLY,
     "blogId=32213",
     [("调用列表接口获取该博客", "记录 title、content、id", "获取成功"),
      ("调用详情接口获取同一博客", "记录 title、content、id", "获取成功"),
      ("比对三个字段", "两接口取值应完全一致", "标题、内容、id 均一致")],
     "通过"),

    ("H1", "H 安全", "标题中的 HTML 与脚本不被执行", "P0", "错误推测",
     LOGIN_PRE,
     "拦截详情接口返回标题：<img src=x onerror=\"window.__dx1=1\">",
     [("拦截详情接口返回含脚本的标题", "拦截生效", "拦截生效"),
      ("打开详情页并等待渲染", "标题被渲染", "渲染完成"),
      ("检查是否产生 img 元素", "标题用 .text() 赋值，不应产生元素", "元素数 0"),
      ("检查脚本是否执行", "不应执行", "window.__dx1 为 null，安全")],
     "通过"),
    ("H2", "H 安全", "正文中的脚本不被执行", "P0", "错误推测",
     LOGIN_PRE,
     "拦截详情接口返回正文含 <img onerror> 与 <script>",
     [("拦截详情接口返回含脚本的正文", "拦截生效", "拦截生效"),
      ("打开详情页并等待 Markdown 渲染", "正文渲染完成", "渲染完成"),
      ("检查两个注入点是否触发", "均不应执行", "两个注入点均未执行，markdownToHTML 处理安全")],
     "通过"),
    ("H3", "H 安全", "作者信息接口泄露密码哈希", "P0", "错误推测",
     LOGIN_PRE + READONLY,
     "GET /user/getAuthorInfo?blogId=32213（带 token）",
     [("以有效 token 调用作者信息接口", "返回 200 且 code=SUCCESS", "返回成功"),
      ("列出响应 data 的所有字段", "应仅返回前端需要的字段", "字段含 id、userName、password、githubUrl 等"),
      ("检查是否包含敏感字段", "不应返回 password", "包含 password 且为密码哈希"),
      ("评估影响", "敏感信息泄露属高危缺陷", "确认泄露，已记为 BUG-D01")],
     "发现缺陷 BUG-D01"),
    ("H4", "H 安全", "更新不存在的 id 被误报为成功", "P1", "错误推测",
     "1.有效 token\n2.不存在的 id",
     "POST /blog/update，id=99999999",
     [("以有效 token 更新不存在的 id", "应返回失败", "返回 code=SUCCESS、data=true"),
      ("评估影响", "误报成功会造成\"已保存\"假象", "确认误报成功，已记为 BUG-D03")],
     "发现缺陷 BUG-D03"),
    ("H5", "H 安全", "缺少 id 的更新不误改数据", "P1", "边界值分析",
     "1.有效 token\n2.不带 id",
     "POST /blog/update，仅带 title 与 content" + READONLY,
     [("记录更新前的博客列表", "作为比对基线", "基线已记录"),
      ("不带 id 提交更新请求", "应拒绝且不修改任何数据", "返回 code=SUCCESS、data=false"),
      ("再次查询列表", "无任何博客被修改，条数不变", "无博客含探测标题，条数不变，WHERE 条件正常")],
     "通过"),
    ("H6", "H 安全", "删除不存在的 id 不误删其他数据", "P1", "错误推测",
     "1.有效 token\n2.不存在的 id",
     "POST /blog/delete?blogId=99999999" + READONLY,
     [("记录删除前的博客列表", "作为比对基线", "基线已记录"),
      ("删除不存在的 id", "应拒绝且不删除任何数据", "返回成功但对不存在记录误报（BUG-D03）"),
      ("再次查询列表", "条数不变，原有博客完好", "条数不变，32213 完好，未误删")],
     "通过"),

    ("I1", "I 交互与兼容", "导航链接跳转正确", "P2", "场景法",
     LOGIN_PRE + READONLY,
     "点击\"主页\"与\"写博客\"",
     [("点击\"主页\"", "跳转 blog_list.html", "跳转成功"),
      ("返回详情页后点击\"写博客\"", "跳转 blog_edit.html", "跳转成功")],
     "通过"),
    ("I2", "I 交互与兼容", "注销清除 token 并跳转登录页", "P1", "场景法",
     LOGIN_PRE + READONLY,
     "点击\"注销\"",
     [("确认已登录（token 存在）", "token 非空", "token 存在"),
      ("点击\"注销\"", "清除 token 并跳转登录页", "token 已清除，跳转登录页")],
     "通过"),
    ("I3", "I 交互与兼容", "窄屏下正文横向溢出（固定宽度布局）", "P2", "兼容性测试",
     LOGIN_PRE + READONLY,
     "分辨率：1920x1080、1366x768、768x1024、375x667",
     [("依次设置四种分辨率并打开详情页", "记录布局表现", "已测量"),
      ("测量 .container 宽度与文档宽度", "容器宽度应随视口收缩", "容器固定 800px，不收缩"),
      ("检查是否出现横向滚动条", "窄屏下不应出现横向滚动", "768px 与 375px 下出现横向溢出"),
      ("评估影响", "固定宽度布局在窄屏下体验差", "确认溢出，已记为 BUG-D07")],
     "发现缺陷 BUG-D07"),
    ("I4", "I 交互与兼容", "静态资源无 4xx/5xx", "P2", "界面检查",
     LOGIN_PRE + READONLY,
     "访问详情页并监听响应",
     [("打开详情页并监听所有响应", "页面 HTML 返回 200", "返回 200"),
      ("记录状态码 >= 400 的资源", "editor.md 等依赖应加载成功", "本次运行全部资源正常")],
     "通过"),
    ("I5", "I 交互与兼容", "无效 token 访问详情页被拦截", "P0", "错误推测",
     "1.向 localStorage 写入伪造 token",
     "伪造 token：garbage.token.xxx",
     [("写入伪造 token", "token 已写入", "写入成功"),
      ("访问详情页", "接口返回 401，跳转登录页", "跳转 blog_login.html，防护有效")],
     "通过"),
]

DEFECTS = [
    ("BUG-D01", "高", "安全", "/user/getAuthorInfo 响应体泄露密码哈希",
     "获取博客作者信息的接口在 data 中直接返回 password 字段，值为用户密码哈希。"
     "详情页仅需 userName 与 githubUrl，完全不需要密码字段。",
     "1）登录后以有效 token 调用 /user/getAuthorInfo?blogId=32213；\n"
     "2）查看响应 data 的全部字段。",
     "响应字段为 [id, userName, password, githubUrl, deleteFlag, createTime, updateTime]，"
     "其中 password 为密码哈希值。",
     "任何登录用户都可通过该接口获取任意博客作者的密码哈希，结合列表页的高危 XSS"
     "（BUG-L01）可批量窃取，风险叠加。此缺陷与列表页 BUG-L02（getUserInfo）为同类问题，"
     "说明接口层普遍未做敏感字段过滤。",
     "统一在接口返回前剔除 password 等敏感字段，改用 VO/DTO 只输出前端需要的字段。",
     "H3", "待修复"),
    ("BUG-D02", "中", "数据一致性", "loginUser 字段在列表与详情接口间矛盾",
     "同一篇博客，列表接口 /blog/getList 返回 loginUser=false，"
     "详情接口 /blog/getBlogDetail 返回 loginUser=true。该字段用于决定是否显示"
     "\"编辑/删除\"按钮，取值矛盾意味着字段语义不可信。",
     "1）以 zhangsan 登录；\n2）调用 /blog/getList 查看博客 32213 的 loginUser；\n"
     "3）调用 /blog/getBlogDetail?blogId=32213 查看同一字段。",
     "列表接口 loginUser=false，详情接口 loginUser=true，两者矛盾。",
     "权限展示所依赖的字段不可信：若某接口把他人博客算成 true，会在前端暴露"
     "编辑/删除入口（后端是否兜底校验未确认，见 BUG-D03）。",
     "统一 loginUser 的计算逻辑并补充自动化断言；权限判断应在后端强制校验，"
     "不能只依赖前端字段。",
     "D3", "待修复"),
    ("BUG-D03", "中", "数据完整性", "更新/删除不存在的 id 被误报为成功",
     "/blog/update 与 /blog/delete 对不存在的 blogId 仍返回 code=SUCCESS、data=true，"
     "成功判定不基于数据库影响行数。",
     "1）以有效 token 提交 POST /blog/update，id=99999999；\n"
     "2）以有效 token 提交 POST /blog/delete?blogId=99999999；\n3）查看响应。",
     "两个接口均返回 {\"code\":\"SUCCESS\",\"data\":true}；实际未修改/删除任何记录"
     "（已验证其他数据未被误改，WHERE 条件正常）。",
     "用户会收到\"操作成功\"的假象，但数据并未变化；同时说明接口未校验记录是否存在，"
     "归属校验也极可能缺失，是越权风险的重要信号。",
     "更新/删除应基于影响行数判定成功（affectedRows > 0），记录不存在时返回明确失败；"
     "并强制校验 blog.userId 是否等于当前登录用户。",
     "D4、H4、H6", "待修复"),
    ("BUG-D04", "中", "可用性", "无效 blogId 错误信息不可诊断且无友好空态",
     "参数非法（abc）、记录不存在（99999999）、参数缺失三种情况返回完全相同的"
     "\"内部错误, 请联系管理员\"；页面仅弹该提示，标题与正文均空白，无返回入口。",
     "1）分别以 blogId=abc、blogId=99999999、不带 blogId 访问详情接口；\n"
     "2）在页面上以不存在的 id 打开详情页。",
     "三种情况接口响应完全相同；页面上标题、日期、正文均为空且无操作按钮，"
     "仅弹出\"内部错误, 请联系管理员\"。",
     "用户与开发者都无法区分是链接错误、文章已删除还是系统故障；"
     "用户停留在空白页面无法继续操作，属可用性与可维护性问题。",
     "区分参数非法与记录不存在并返回明确 errMsg；页面在无数据时展示空态提示"
     "（如\"该博客不存在或已被删除\"）并提供返回列表入口。",
     "C1、C2、C3", "待修复"),
    ("BUG-D05", "低", "数据一致性", "左侧文章计数为硬编码，与实际不符",
     "左侧卡片中\"文章/分类\"的计数为 HTML 静态写死（2 和 1），未随实际数据变化。",
     "1）调用 /blog/getList 统计当前账号实际博客数；\n"
     "2）打开详情页读取卡片显示的文章数；\n3）比对两者。",
     "卡片显示文章数=2，实际博客数=1，偏差 1。",
     "用户看到的统计与实际内容不符，误导用户。该问题在列表页同样存在（BUG-L03），"
     "属系统性硬编码问题。",
     "计数改为动态获取（由接口返回统计数据或按列表长度计算）。",
     "G1", "待修复"),
    ("BUG-D06", "低", "代码质量", "详情页存在重复与陈旧代码",
     "页面源码中保留了一个无参的 deleteBlog() 函数（且带 confirm 逻辑）、"
     "重复定义的 userUrl 变量，以及被注释掉的 getUserInfo 调用。",
     "1）查看 /blog_detail.html 源码；\n2）检查脚本段的函数与变量定义。",
     "存在被注释的 deleteBlog() 定义与重复的 userUrl 变量声明；"
     "editormd.markdownToHTML 与页面注释中的说明也不一致。",
     "死代码增加维护成本与误改风险，重复声明的变量可能导致后续修改时行为不符预期。",
     "清理注释掉的死代码与重复声明，统一由后端返回值驱动前端逻辑。",
     "—", "待修复"),
    ("BUG-D07", "低", "兼容性", "窄屏下出现横向滚动条",
     ".container 为固定宽度 800px，未做响应式收缩。在 768px、375px 等窄视口下，"
     "页面宽度超出视口，出现横向滚动条。",
     "1）打开详情页；\n2）将视口依次设为 1920、1366、768、375 宽度；\n"
     "3）测量 container 宽度与 document.scrollWidth。",
     "container 宽度恒为 800px；视口 768px 时文档宽 800px（溢出 32px），"
     "视口 375px 时文档宽 800px（溢出 425px），均出现横向滚动。",
     "移动端或窄窗口下用户需要左右拖动才能阅读正文，体验较差。"
     "该问题同样存在于列表页与编辑页（共用 common.css）。",
     "为 .container 设置 max-width: 100% 并配合媒体查询做响应式布局。",
     "I3", "待修复"),
]

MODULES = [
    ("A 访问控制", "验证未登录访问与越权写入的防护", "安全测试"),
    ("B 数据渲染", "验证标题、日期、正文与作者信息的渲染", "功能测试"),
    ("C 无效与异常 blogId", "验证异常参数下的错误处理与提示", "健壮性测试"),
    ("D 权限与操作按钮", "验证编辑/删除按钮的权限控制与归属校验", "安全测试"),
    ("E 编辑跳转", "验证编辑按钮跳转与更新页数据预填", "功能测试"),
    ("F 删除流程", "验证删除确认、取消与真实删除的完整流程", "功能测试"),
    ("G 数据一致性", "验证展示数据与实际数据的一致性", "功能测试"),
    ("H 安全", "验证 XSS 防护与敏感信息泄露", "安全测试"),
    ("I 交互与兼容", "验证导航、注销、多分辨率与静态资源", "兼容性测试"),
]

wb = Workbook()

# ===== Sheet1 用例说明 =====
ws0 = wb.active
ws0.title = "用例说明"
ws0.sheet_view.showGridLines = False
ws0["A1"] = "博客详情页 测试用例"
ws0["A1"].font = TITLE_FONT
ws0.merge_cells("A1:B1")
rows0 = [
    ("被测系统", "博客系统（我的博客系统）"),
    ("被测页面", "http://127.0.0.1:9090/blog_detail.html?blogId=32213（页面标题：博客详情页）"),
    ("核心功能", "展示博客标题/日期/正文与作者信息；按权限渲染\"编辑/删除\"按钮"),
    ("关键接口", "GET /blog/getBlogDetail（详情）\nGET /user/getAuthorInfo（作者信息）\n"
                 "POST /blog/delete（删除）\nPOST /blog/update（更新，由编辑按钮跳转的更新页调用）"),
    ("前置依赖", "需登录态；localStorage.user_token 由 /user/login 获取"),
    ("测试账号", "zhangsan / 123456"),
    ("用例总数", f"{len(CASES)} 条"),
    ("测试步骤总数", f"{sum(len(c[7]) for c in CASES)} 步"),
    ("执行结果", "全部执行通过；其中 9 条用例在执行中发现缺陷（详见\"缺陷清单\"）"),
    ("", ""),
    ("【数据安全策略】", ""),
    ("只读用例", "使用环境中原有博客 32213，仅读取不修改"),
    ("写操作用例", "只针对**自建测试博客**（标题含 __E2E_DETAIL__ 标记）执行删除与更新验证"),
    ("强制清理", "每个涉及自建数据的用例都在 finally 中删除测试博客，并断言环境恢复；"
                 "测试结束后环境仅剩原有博客 32213"),
    ("零风险验证", "越权与异常用例使用不存在的 id（如 99999999），不触碰任何真实数据"),
    ("", ""),
    ("【未能结论的风险项】", ""),
    ("跨用户越权", "环境无注册接口，无法创建第二个用户，因此**未能实证**\"用户 A 能否改删用户 B 的博客\"。"
                   "已确认更新/删除接口对不存在的 id 误报成功（BUG-D03），故归属校验存在风险，"
                   "建议后端补充 blog.userId 校验并单独验证。"),
    ("", ""),
    ("【字段说明】", ""),
    ("用例编号", "模块字母 + 序号，如 F3 表示 F 模块第 3 条用例"),
    ("所属模块", "按测试类型划分为 A~I 共 9 个模块"),
    ("优先级", "P0 核心必测；P1 重要；P2 一般"),
    ("设计方法", "等价类划分、边界值分析、错误推测、场景法、界面检查、兼容性测试等"),
    ("前置条件", "执行该用例前必须满足的环境与登录态"),
    ("测试数据", "该用例使用的具体输入；标注【只读】/【数据安全】的为受控数据"),
    ("步骤序号", "用例内的步骤序号，每条用例从 1 开始连续编号"),
    ("测试步骤 / 预期结果 / 实际结果", "该步骤的操作、应有结果与实际观测结果"),
    ("结论", "该用例整体执行结论"),
]
r = 3
for k, v in rows0:
    c1 = ws0.cell(r, 1, k)
    c1.font = BOLD if not k.startswith("【") else Font(bold=True, size=11, name="微软雅黑", color="C00000")
    c1.fill = SEC_FILL
    c1.border = BORDER
    c1.alignment = WRAP
    c2 = ws0.cell(r, 2, v)
    c2.font = FONT
    c2.border = BORDER
    c2.alignment = WRAP
    ws0.row_dimensions[r].height = 20 if len(v) < 60 else 52
    r += 1
ws0.column_dimensions["A"].width = 20
ws0.column_dimensions["B"].width = 94

# ===== Sheet2 测试用例 =====
ws = wb.create_sheet("测试用例")
ws.sheet_view.showGridLines = False
ws.page_setup.orientation = "landscape"
ws.sheet_properties.pageSetUpPr.fitToPage = True
ws.page_setup.fitToWidth = 1
HEADERS = ["用例编号", "所属模块", "用例标题", "优先级", "设计方法",
           "前置条件", "测试数据", "步骤序号", "测试步骤", "预期结果", "实际结果", "结论"]
WIDTHS = [9, 16, 30, 7, 12, 30, 30, 7, 36, 32, 38, 14]
for i, (h, w) in enumerate(zip(HEADERS, WIDTHS), start=1):
    c = ws.cell(1, i, h)
    c.font = HDR_FONT; c.fill = HDR_FILL; c.alignment = CENTER; c.border = BORDER
    ws.column_dimensions[get_column_letter(i)].width = w
ws.row_dimensions[1].height = 26
ws.freeze_panes = "A2"
ws.auto_filter.ref = f"A1:L{1 + sum(len(c[7]) for c in CASES)}"
ST_COLOR = {"发现缺陷": "C00000", "通过": "1F6B33"}
row = 2
for idx, case in enumerate(CASES):
    cid, mod, case_title, pri, method, pre, data, steps, result = case
    fill = CASE_A if idx % 2 == 0 else CASE_B
    start = row
    for si, (op, exp, act) in enumerate(steps, start=1):
        vals = [cid, mod, case_title, pri, method, pre, data, si, op, exp, act, ""]
        for ci, v in enumerate(vals, start=1):
            cell = ws.cell(row, ci, v)
            cell.font = FONT; cell.border = BORDER; cell.fill = fill
            cell.alignment = CENTER if ci in (4, 8) else WRAP
        if pri == "P0":
            ws.cell(row, 4).font = BOLD_RED
        elif pri == "P1":
            ws.cell(row, 4).font = BOLD
        if si == len(steps):
            cc = ws.cell(row, 12, result)
            cc.font = Font(bold=True, size=10, name="微软雅黑",
                           color=ST_COLOR["发现缺陷"] if result.startswith("发现缺陷") else ST_COLOR["通过"])
            cc.alignment = WRAP_C; cc.border = BORDER; cc.fill = fill
        row += 1
    for ci in (1, 2, 3, 4, 5, 6, 7):
        if len(steps) > 1:
            ws.merge_cells(start_row=start, start_column=ci, end_row=row - 1, end_column=ci)
        ws.cell(start, ci).alignment = Alignment(
            wrap_text=True, vertical="center",
            horizontal="center" if ci in (1, 2, 4, 5) else "left")

# ===== Sheet3 模块划分 =====
ws1 = wb.create_sheet("模块划分")
ws1.sheet_view.showGridLines = False
ws1["A1"] = "测试模块划分"
ws1["A1"].font = TITLE_FONT
ws1.merge_cells("A1:F1")
for i, (h, w) in enumerate(zip(["模块", "模块说明", "测试类型", "用例数", "步骤数", "用例编号"],
                               [22, 46, 14, 9, 9, 22]), start=1):
    c = ws1.cell(3, i, h)
    c.font = HDR_FONT; c.fill = HDR_FILL; c.alignment = CENTER; c.border = BORDER
    ws1.column_dimensions[get_column_letter(i)].width = w
r = 4
for name, desc, ttype in MODULES:
    sub = [c for c in CASES if c[1] == name]
    vals = [name, desc, ttype, len(sub), sum(len(c[7]) for c in sub),
            "、".join(c[0] for c in sub)]
    for ci, v in enumerate(vals, start=1):
        cell = ws1.cell(r, ci, v)
        cell.font = FONT; cell.border = BORDER
        cell.alignment = WRAP if ci == 2 else WRAP_C
    r += 1
tot = ["合计", "", "", len(CASES), sum(len(c[7]) for c in CASES),
       f"{len(CASES)} 条用例 / {sum(len(c[7]) for c in CASES)} 个步骤"]
for ci, v in enumerate(tot, start=1):
    cell = ws1.cell(r, ci, v)
    cell.font = BOLD; cell.fill = SEC_FILL; cell.border = BORDER; cell.alignment = WRAP_C

# ===== Sheet4 缺陷清单 =====
ws2 = wb.create_sheet("缺陷清单")
ws2.sheet_view.showGridLines = False
ws2.page_setup.orientation = "landscape"
ws2.sheet_properties.pageSetUpPr.fitToPage = True
ws2.page_setup.fitToWidth = 1
dh = ["缺陷编号", "严重级别", "类别", "缺陷标题", "缺陷描述",
      "复现步骤", "实际结果", "影响", "修复建议", "关联用例", "状态"]
dw = [11, 10, 12, 30, 40, 38, 40, 36, 34, 12, 10]
for i, (h, w) in enumerate(zip(dh, dw), start=1):
    c = ws2.cell(1, i, h)
    c.font = HDR_FONT; c.fill = HDR_FILL; c.alignment = CENTER; c.border = BORDER
    ws2.column_dimensions[get_column_letter(i)].width = w
ws2.row_dimensions[1].height = 26
ws2.freeze_panes = "A2"
SEV_FILL = {"高": PatternFill("solid", fgColor="F4B6B6"),
            "中": PatternFill("solid", fgColor="FDE9C9"),
            "低": PatternFill("solid", fgColor="E2EFDA")}
for ri, d in enumerate(DEFECTS, start=2):
    for ci, v in enumerate(d, start=1):
        c = ws2.cell(ri, ci, v)
        c.font = FONT; c.border = BORDER
        c.alignment = WRAP_C if ci in (1, 2, 3, 10, 11) else WRAP
    ws2.cell(ri, 2).fill = SEV_FILL[d[1]]
    ws2.cell(ri, 2).font = BOLD

# ===== Sheet5 执行统计 =====
ws3 = wb.create_sheet("执行统计")
ws3.sheet_view.showGridLines = False
ws3["A1"] = "执行统计"
ws3["A1"].font = TITLE_FONT
ws3.merge_cells("A1:C1")
for i, (h, w) in enumerate(zip(["统计项", "数值", "说明"], [24, 16, 58]), start=1):
    c = ws3.cell(3, i, h)
    c.font = HDR_FONT; c.fill = HDR_FILL; c.alignment = CENTER; c.border = BORDER
    ws3.column_dimensions[get_column_letter(i)].width = w
defect_cases = [c[0] for c in CASES if c[8].startswith("发现缺陷")]
stats = [
    ("用例总数", f"{len(CASES)} 条", "见\"测试用例\"工作表"),
    ("测试步骤总数", f"{sum(len(c[7]) for c in CASES)} 步", "每条用例按步骤逐行展开"),
    ("自动化执行次数", f"{len(CASES)} 次", "1 条用例 = 1 次执行"),
    ("执行通过", f"{len(CASES)} 次", "断言层面无失败"),
    ("通过率", "100%", f"{len(CASES)} / {len(CASES)}"),
    ("发现缺陷", f"{len(DEFECTS)} 个", "1 高 / 3 中 / 3 低，详见\"缺陷清单\""),
    ("结论为\"发现缺陷\"的用例", f"{len(defect_cases)} 条", "、".join(defect_cases)),
    ("数据安全", "无残留", "测试全程只用自建测试博客，结束后环境仅剩原有博客 32213"),
    ("未能结论项", "1 项", "跨用户越权（环境无注册接口，无法创建第二用户，未能实证）"),
]
r = 4
for k, v, note in stats:
    for ci, val in enumerate([k, v, note], start=1):
        c = ws3.cell(r, ci, val)
        c.font = BOLD if ci == 1 else FONT
        c.border = BORDER
        c.alignment = WRAP_C if ci != 3 else WRAP
    r += 1
r += 1
ws3.cell(r, 1, "优先级分布").font = BOLD
r += 1
for ci, h in enumerate(["优先级", "用例数", "含义"], start=1):
    c = ws3.cell(r, ci, h)
    c.font = HDR_FONT; c.fill = HDR_FILL; c.alignment = CENTER; c.border = BORDER
r += 1
for p, mean in [("P0", "核心必测"), ("P1", "重要功能"), ("P2", "一般功能与兼容")]:
    n = sum(1 for c in CASES if c[3] == p)
    for ci, v in enumerate([p, n, mean], start=1):
        c = ws3.cell(r, ci, v)
        c.font = FONT; c.border = BORDER
        c.alignment = WRAP_C if ci != 3 else WRAP
    r += 1

wb.save(OUT)
print("OK cases=%d steps=%d defects=%d defect_cases=%d" % (
    len(CASES), sum(len(c[7]) for c in CASES), len(DEFECTS), len(defect_cases)))
